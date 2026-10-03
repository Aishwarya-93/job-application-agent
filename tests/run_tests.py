"""Test scenarios for the agent and the web UI (run from the project folder:  python tests/run_tests.py).

They use the offline stand-in model so they run without an API key. This checks the tools, memory,
tracing and error handling. It does NOT measure the quality of a real LLM's reasoning.
"""
import os
import subprocess
import sys
import uuid
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
os.chdir(ROOT); sys.path.insert(0, str(ROOT))

from agent import build_agent, run_turn                    # noqa: E402
from agent_trace import TraceHandler                        # noqa: E402
from offline_llm import ScriptedChatModel                   # noqa: E402

LOG = "logs/test_trace.txt"
Path(LOG).unlink(missing_ok=True)
REPORT = Path("outputs/improvement_report.md")
results = []


def session():
    tracer = TraceHandler(LOG, header="TEST RUN (offline stand-in)")
    cfg = {"configurable": {"thread_id": str(uuid.uuid4())}, "callbacks": [tracer], "recursion_limit": 25}
    return build_agent(ScriptedChatModel()), tracer, cfg


def goal(resume, jd):
    return f"Analyze my resume at '{resume}' against the job description at '{jd}'."


def check(name, ok, detail=""):
    results.append((name, ok, detail)); print(f"[{'PASS' if ok else 'FAIL'}] {name} {detail}\n")


# 1 ---------------------------------------------------------------- normal analysis
REPORT.unlink(missing_ok=True)
agent, tr, cfg = session()
ans = run_turn(agent, goal("data/sample_resume.txt", "data/sample_jd.txt"), cfg, tr, first=True)
check("1. Normal analysis (4 tools, in order, report saved, score 70%)",
      tr.calls == ["load_documents", "analyze_skill_gap", "calculate_match_score", "generate_report"]
      and REPORT.exists() and "70%" in ans, f"tools={tr.calls}")

# 3 ---------------------------------------------------------------- follow-up uses memory
ans2 = run_turn(agent, "What are my three biggest skill gaps?", cfg, tr)
check("3. Follow-up answered from memory (no tools re-run)",
      tr.calls == [] and all(s in ans2 for s in ("Data Visualization", "Statistics", "Docker")), f"tools={tr.calls}")
ans3 = run_turn(agent, "What was my score?", cfg, tr)
check("3b. Second follow-up recalls score", tr.calls == [] and "70%" in ans3)

# 2 ---------------------------------------------------------------- different JD, many missing skills
agent, tr, cfg = session()
ans = run_turn(agent, goal("data/sample_resume.txt", "data/sample_jd_devops.txt"), cfg, tr, first=True)
check("2. Missing skills detected on DevOps JD (score 25%)",
      all(s in ans for s in ("Docker", "Kubernetes", "Terraform", "AWS")) and "25%" in ans,
      f"tools={tr.calls}")

# 4 ---------------------------------------------------------------- invalid input
REPORT.unlink(missing_ok=True)
agent, tr, cfg = session()
ans = run_turn(agent, goal("data/NOPE.txt", "data/sample_jd.txt"), cfg, tr, first=True)
check("4a. Missing resume file -> clear error, no report, no crash",
      "not found" in ans and not REPORT.exists() and tr.calls == ["load_documents"], f"tools={tr.calls}")
cp = subprocess.run([sys.executable, "agent.py", "--offline", "--resume", "nope.txt", "--log", LOG],
                    capture_output=True, text=True)
check("4b. CLI with missing file -> exit code != 0, message, no traceback",
      cp.returncode != 0 and "not found" in (cp.stderr + cp.stdout) and "Traceback" not in cp.stderr)
empty = Path("data/_empty.txt"); empty.write_text("")
agent, tr, cfg = session()
ans = run_turn(agent, goal(str(empty), "data/sample_jd.txt"), cfg, tr, first=True)
empty.unlink()
check("4c. Empty resume file -> clear error", "no readable text" in ans)
agent, tr, cfg = session()
ans = run_turn(agent, "Please analyze my resume.", cfg, tr, first=True)
check("4d. No files given -> agent asks for them, calls no tools", tr.calls == [] and "resume file" in ans)

# 5 ---------------------------------------------------------------- tool failure
os.environ["SIMULATE_TOOL_FAILURE"] = "analyze_skill_gap"
agent, tr, cfg = session()
ans = run_turn(agent, goal("data/sample_resume.txt", "data/sample_jd.txt"), cfg, tr, first=True)
del os.environ["SIMULATE_TOOL_FAILURE"]
check("5. Tool failure handled gracefully (error explained, run stops cleanly)",
      "could not finish" in ans and "analyze_skill_gap failed" in ans
      and tr.calls == ["load_documents", "analyze_skill_gap"], f"tools={tr.calls}")

# 6 ---------------------------------------------------------------- tools called out of order
import tools                                                          # noqa: E402
tools.SESSION.clear()
out = tools.calculate_match_score.invoke({})
check("6. Tool called before prerequisites -> error message, not exception", '"status": "error"' in out)

# 7 ---------------------------------------------------------------- PDF resume (optional)
try:
    from reportlab.pdfgen import canvas
    c = canvas.Canvas("data/_resume.pdf")
    for i, line in enumerate(["SKILLS", "Python, SQL, Docker, Statistics, Git", "PROJECTS",
                              "- Built a Docker based Python app"]):
        c.drawString(72, 750 - 20 * i, line)
    c.save()
    agent, tr, cfg = session()
    ans = run_turn(agent, goal("data/_resume.pdf", "data/sample_jd.txt"), cfg, tr, first=True)
    Path("data/_resume.pdf").unlink()
    check("7. PDF resume is read correctly", "Docker" in ans and tr.calls[-1] == "generate_report")
except ImportError:
    print("[SKIP] 7. PDF test (pip install reportlab to enable)\n")

# 8 ---------------------------------------------------------------- live mode without key
env = {k: v for k, v in os.environ.items() if k != "OPENROUTER_API_KEY"}
cp = subprocess.run([sys.executable, "agent.py", "--log", LOG], capture_output=True, text=True, env=env, cwd=ROOT)
check("8. Live mode without API key -> friendly message, no traceback",
      cp.returncode != 0 and "OPENROUTER_API_KEY" in cp.stderr and "Traceback" not in cp.stderr)

# =====================================================================================================
# Web UI / dynamic-input tests (Flask test client, offline stand-in model)
# =====================================================================================================
import glob                                                           # noqa: E402
import io                                                             # noqa: E402
import tempfile                                                       # noqa: E402
import app as webapp                                                  # noqa: E402

webapp.app.config["LLM_FACTORY"] = ScriptedChatModel
webapp.app.config["OFFLINE"] = True
webapp.app.config["TESTING"] = True


def client():
    return webapp.app.test_client()                                   # own cookie jar = own browser session


def upload(c, resume=None, jd_file=None, jd_text=None, sample=False):
    data = {}
    if sample:
        data["use_sample"] = "1"
    if resume:
        data["resume"] = (io.BytesIO(resume[1]), resume[0])
    if jd_file:
        data["jd_file"] = (io.BytesIO(jd_file[1]), jd_file[0])
    if jd_text is not None:
        data["jd_text"] = jd_text
    r = c.post("/analyze", data=data, content_type="multipart/form-data")
    return r.status_code, r.get_json()


DEVOPS_CV = (b"ARUN SHAH\nSKILLS\nDocker, Kubernetes, Terraform, AWS, Linux, Git, Jenkins, Python\n"
             b"PROJECTS\n- Built CI/CD pipeline with Jenkins and Docker on AWS\n")
DEVOPS_JD = (b"Job Title: DevOps Engineer\nRequired Skills:\n- Docker\n- Kubernetes\n- Terraform\n- AWS\n"
             b"- Linux\n- CI/CD\n")
WEB_CV = b"MEERA PATEL\nSKILLS\nJavaScript, React, HTML, CSS, Node.js, Git\nPROJECTS\n- Shop UI built with React and CSS\n"
WEB_JD_TEXT = "Frontend Developer\nRequired Skills:\n- JavaScript\n- React\n- HTML\n- CSS\n- Node.js\n- Docker\n"

# 9 ---------------------------------------------------------------- UI: txt resume + pasted JD, full result
c = client()
code, j = upload(c, ("cv.txt", WEB_CV), jd_text=WEB_JD_TEXT)
r = j.get("results") or {}
check("9. UI: TXT resume + pasted JD -> analysis (score 83%, missing Docker)",
      code == 200 and r.get("score") == 83 and r.get("missing_skills") == ["Docker"]
      and "React" in r.get("matched_skills", []), f"status={code} score={r.get('score')}")
tools_used = [e["tool"] for e in j["trace"] if e["kind"] == "tool_call"]
labels = [e["label"] for e in j["trace"] if e["kind"] in ("goal", "observation", "final")]
check("9b. UI: execution trace has goal -> 4 tool steps -> final, with tool names",
      tools_used == ["load_documents", "analyze_skill_gap", "calculate_match_score", "generate_report"]
      and labels[0] == "Goal received" and labels[-1] == "Final response prepared"
      and "Report generated" in labels, f"{labels}")
check("9c. UI: recommendations, learning areas and summary come back for display",
      r.get("recommendations") and r.get("learning_areas") and r.get("summary") and r.get("report_ready"))

# 10 --------------------------------------------------------------- UI: follow-up uses memory of THAT upload
fr = c.post("/ask", json={"question": "What are my three biggest skill gaps?"}).get_json()
check("10. UI follow-up answered from memory about the uploaded files (no tools re-run)",
      fr["ok"] and "Docker" in fr["answer"] and not [e for e in fr["trace"] if e["kind"] == "tool_call"]
      and any(e["kind"] == "memory" for e in fr["trace"]), fr.get("answer", "")[:90])
check("10b. Follow-up before any analysis -> clear message",
      client().post("/ask", json={"question": "hi"}).status_code == 400)

# 11 --------------------------------------------------------------- report is generated from the upload + downloadable
dl = c.get("/download")
check("11. Report downloads and contains the uploaded resume's content, not sample data",
      dl.status_code == 200 and b"Shop UI built with React" in dl.data and b"Student Performance" not in dl.data
      and b"83%" in dl.data)

# 12 --------------------------------------------------------------- CHANGING the inputs changes the output
c2 = client()
code2, j2 = upload(c2, ("devops.txt", DEVOPS_CV), jd_file=("devops_jd.txt", DEVOPS_JD))
r2 = j2["results"]
check("12. Different resume + JD -> different output (DevOps: 100%, nothing missing)",
      code2 == 200 and r2["score"] == 100 and r2["missing_skills"] == [] and r2["score"] != r["score"]
      and "Kubernetes" in r2["matched_skills"], f"score={r2['score']}")
code3, j3 = upload(c2, ("cv.txt", WEB_CV), jd_file=("devops_jd.txt", DEVOPS_JD))
r3 = j3["results"]
check("12b. Same JD, other resume -> score changes (web resume vs DevOps JD: 0%, 5+ missing)",
      r3["score"] == 0 and len(r3["missing_skills"]) == 6, f"score={r3['score']}")
code4, j4 = upload(c2, sample=True)
check("12c. Demo fallback still works (sample files: 70%)", code4 == 200 and j4["results"]["score"] == 70)
rep = c2.get("/download").data
check("12d. New analysis replaces the old report (no stale content)",
      b"70%" in rep and b"Kubernetes" not in rep.split(b"## Missing Skills")[0])

# 13 --------------------------------------------------------------- DOCX and PDF uploads
import docx                                                           # noqa: E402
buf = io.BytesIO(); d = docx.Document()
d.add_paragraph("PRIYA NAIR"); d.add_paragraph("SKILLS"); d.add_paragraph("Python, SQL, Docker, Statistics, Git, Tableau")
d.add_paragraph("PROJECTS"); d.add_paragraph("- Dashboard in Tableau with Python and SQL")
t = d.add_table(rows=1, cols=2); t.rows[0].cells[0].text = "Cloud"; t.rows[0].cells[1].text = "AWS"
d.save(buf)
code, j = upload(client(), ("resume.docx", buf.getvalue()), jd_file=("jd.txt", open("data/sample_jd.txt", "rb").read()))
rr = j.get("results") or {}
check("13. DOCX resume (incl. table text) is read: Docker, Statistics, Data Visualization matched",
      code == 200 and not {"Docker", "Statistics", "Data Visualization"} & set(rr.get("missing_skills", ["x"])),
      f"missing={rr.get('missing_skills')}")
try:
    from reportlab.pdfgen import canvas
    pb = io.BytesIO(); cv = canvas.Canvas(pb)
    for i, line in enumerate(["SKILLS", "Python, SQL, Docker, Statistics, Git", "PROJECTS", "- Built a Docker based Python app"]):
        cv.drawString(72, 750 - 20 * i, line)
    cv.save()
    code, j = upload(client(), ("resume.pdf", pb.getvalue()), jd_text=open("data/sample_jd.txt").read())
    rr = j.get("results") or {}
    check("13b. PDF resume + pasted JD works", code == 200 and "Docker" in rr.get("matched_skills", []))
except ImportError:
    print("[SKIP] 13b. PDF upload test (pip install reportlab to enable)\n")

# 14 --------------------------------------------------------------- missing / bad inputs
c = client()
code, j = upload(c, jd_text="Python")
check("14a. No resume -> 400 with message", code == 400 and "resume" in j["error"].lower())
code, j = upload(c, resume=("cv.txt", WEB_CV))
check("14b. No JD -> 400 with message", code == 400 and "job description" in j["error"].lower())
code, j = upload(c, resume=("cv.txt", WEB_CV), jd_file=("jd.txt", DEVOPS_JD), jd_text="also pasted")
check("14c. JD as file AND text -> 400 (ambiguous)", code == 400 and "not both" in j["error"])
code, j = upload(c, resume=("cv.exe", b"x"), jd_text="Python")
check("14d. Unsupported file type -> 400", code == 400 and "unsupported" in j["error"].lower())
code, j = upload(c, resume=("cv.txt", b"   "), jd_text="Python")
check("14e. Empty resume file -> 400, no agent run", code == 400 and "no readable text" in j["error"])
code, j = upload(c, resume=("cv.pdf", b"not really a pdf"), jd_text="Python")
check("14f. Corrupt PDF -> 400 friendly message", code == 400 and "Could not read" in j["error"])
code, j = upload(c, resume=("cv.txt", WEB_CV), jd_text="We want someone friendly and hardworking.")
check("14g. JD with no recognisable skills -> agent reports problem, no fake score",
      j["ok"] is False and j["results"] is None and "no recognisable required skills" in j["answer"], j.get("error", "")[:80])

# 15 --------------------------------------------------------------- live mode without key; privacy of uploads
webapp.app.config["LLM_FACTORY"], webapp.app.config["OFFLINE"] = None, False
saved_key = os.environ.pop("OPENROUTER_API_KEY", None)
code, j = upload(client(), resume=("cv.txt", WEB_CV), jd_text=WEB_JD_TEXT)
if saved_key:
    os.environ["OPENROUTER_API_KEY"] = saved_key
webapp.app.config["LLM_FACTORY"], webapp.app.config["OFFLINE"] = ScriptedChatModel, True
check("15. Live mode without API key -> 503 with setup message (no traceback)",
      code == 503 and "OPENROUTER_API_KEY" in j["error"])
check("15b. Uploaded files are deleted after the run", not glob.glob(os.path.join(tempfile.gettempdir(), "jobassist_*")))

# 16 --------------------------------------------------------------- nothing hard-coded to the sample files
src = "".join(Path(f).read_text(encoding="utf-8") for f in ("tools.py", "offline_llm.py"))
check("16. tools.py / offline_llm.py never reference the sample files", "sample_resume" not in src and "sample_jd" not in src)

# ---------------------------------------------------------------- summary
print("=" * 60)
for n, ok, _ in results:
    print(("PASS  " if ok else "FAIL  ") + n)
passed = sum(ok for _, ok, _ in results)
print(f"\n{passed}/{len(results)} checks passed")
Path("logs/test_results.txt").write_text("\n".join(("PASS  " if ok else "FAIL  ") + n for n, ok, _ in results)
                                         + f"\n{passed}/{len(results)} checks passed\n")
sys.exit(0 if passed == len(results) else 1)
