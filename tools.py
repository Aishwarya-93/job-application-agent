"""Tools for the Job Application Assistant agent (4 tools).

    1. load_documents       - reads the resume + job description (.txt/.md/.pdf/.docx) ONCE
    2. analyze_skill_gap    - required / matched / missing skills + relevant experience
    3. calculate_match_score- Resume-JD Skill Match Score (simple, explainable)
    4. generate_report      - structured improvement report (markdown)

Working memory
--------------
`SESSION` is the agent's working memory. load_documents stores the parsed documents in it and
later tools read from it, so no tool re-reads a file and the LLM never has to copy long text
between tool calls. (Conversation memory is separate: see MemorySaver in agent.py.)

Every tool returns a JSON string with a short human-readable "summary" - that is what the
execution trace shows - and never raises: failures come back as {"status": "error", ...} so the
agent can explain the problem instead of crashing.
"""
import functools
import json
import os
import re
from pathlib import Path

from langchain_core.tools import tool

SESSION: dict = {}  # working memory: resume, jd, gap, score

# ------------------------------------------------------------------ skill lexicon
# canonical skill -> aliases (lowercase). Deliberately small and easy to extend.
SKILLS = {
    "Python": ["python"], "Java": ["java"], "JavaScript": ["javascript"], "C++": ["c++"],
    "SQL": ["sql", "mysql", "postgresql", "sqlite"],
    "Pandas": ["pandas"], "NumPy": ["numpy"],
    "Scikit-learn": ["scikit-learn", "sklearn", "scikit learn"],
    "Machine Learning": ["machine learning"], "Deep Learning": ["deep learning"],
    "TensorFlow": ["tensorflow"], "PyTorch": ["pytorch"], "NLP": ["nlp", "natural language processing"],
    "Data Visualization": ["data visualization", "data visualisation", "matplotlib", "seaborn",
                           "tableau", "power bi", "plotly"],
    "Power BI": ["power bi", "powerbi"],
    "Statistics": ["statistics", "statistical"], "Excel": ["excel"],
    "Docker": ["docker"], "Kubernetes": ["kubernetes", "k8s"], "Git": ["git", "github"],
    "CI/CD": ["ci/cd", "continuous integration"], "Jenkins": ["jenkins"], "Terraform": ["terraform"],
    "AWS": ["aws", "amazon web services"], "Azure": ["azure"], "GCP": ["gcp", "google cloud"],
    "Linux": ["linux"], "REST APIs": ["rest api", "rest apis", "restful"],
    "React": ["react"], "Node.js": ["node.js", "nodejs"], "HTML": ["html"], "CSS": ["css"],
    "Flask": ["flask"], "FastAPI": ["fastapi"], "Spark": ["spark", "pyspark"],
}
PREFERRED_HEADING = re.compile(r"^\s*(preferred|nice to have|good to have|bonus)", re.I)
EXPERIENCE_SECTIONS = ("PROJECT", "INTERNSHIP", "EXPERIENCE", "WORK")


def _find_skills(text: str) -> dict:
    """Return {canonical skill: position of first mention} for every lexicon skill in text."""
    low, found = text.lower(), {}
    for canon, aliases in SKILLS.items():
        hits = [m.start() for a in aliases
                for m in [re.search(rf"(?<![\w+#.]){re.escape(a)}(?![\w+#])", low)] if m]
        if hits:
            found[canon] = min(hits)
    return found


def _ok(summary: str, **data) -> str:
    return json.dumps({"status": "ok", "summary": summary, **data})


def _err(message: str) -> str:
    return json.dumps({"status": "error", "summary": f"ERROR: {message}"})


def safe_tool(fn):
    """Turn any exception into a clean error observation (and allow a test failure hook)."""
    @functools.wraps(fn)
    def wrapper(*args, **kwargs):
        try:
            if os.getenv("SIMULATE_TOOL_FAILURE") == fn.__name__:  # used only by tests/run_tests.py
                raise RuntimeError("simulated tool failure")
            return fn(*args, **kwargs)
        except Exception as e:  # noqa: BLE001 - we deliberately report, not crash
            return _err(f"{fn.__name__} failed: {e}")
    return wrapper


SUPPORTED_EXTENSIONS = (".txt", ".md", ".pdf", ".docx")


def read_document(path: str, label: str = "Document") -> str:
    """Read .txt / .md / .pdf / .docx into plain text. Raises FileNotFoundError / ValueError."""
    p = Path(path.strip().strip("'\""))
    if not p.is_file():
        raise FileNotFoundError(f"{label} file not found: {path}")
    ext = p.suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader
        text = "\n".join(pg.extract_text() or "" for pg in PdfReader(str(p)).pages)
    elif ext == ".docx":
        import docx
        d = docx.Document(str(p))
        parts = [para.text for para in d.paragraphs]
        for table in d.tables:                      # skills are often kept in tables
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        text = "\n".join(parts)
    elif ext in (".txt", ".md"):
        text = p.read_text(encoding="utf-8", errors="ignore")
    else:
        raise ValueError(f"{label} file type '{p.suffix}' is not supported (use .txt, .pdf or .docx)")
    if not text.strip():
        raise ValueError(f"{label} file has no readable text (empty or scanned PDF): {path}")
    return text


_read_file = read_document  # backwards-compatible name


# ------------------------------------------------------------------------- tools
@tool
@safe_tool
def load_documents(resume_path: str, jd_path: str) -> str:
    """Read the resume and the job description (.txt, .pdf or .docx file paths) and keep them in the
    agent's working memory. Call this FIRST. Other tools take no document arguments."""
    problems, texts = [], {}
    for label, path in (("Resume", resume_path), ("Job description", jd_path)):
        try:
            texts[label] = _read_file(path, label)
        except (FileNotFoundError, ValueError) as e:
            problems.append(str(e))
    if problems:
        return _err("; ".join(problems))
    SESSION.clear()
    SESSION.update(resume=texts["Resume"], jd=texts["Job description"],
                   resume_path=resume_path, jd_path=jd_path)
    r, j = len(texts["Resume"].split()), len(texts["Job description"].split())
    return _ok(f"Resume ({r} words) and job description ({j} words) loaded into working memory.",
               job_title=texts["Job description"].strip().splitlines()[0][:80])


@tool
@safe_tool
def analyze_skill_gap() -> str:
    """Analyse the loaded resume against the loaded job description. Returns required skills,
    matched skills, missing skills (major gaps, in the order the JD lists them), preferred
    skills, and the resume lines that show relevant experience. Needs load_documents first."""
    if "resume" not in SESSION:
        return _err("no documents loaded yet - call load_documents first")
    jd_lines = SESSION["jd"].splitlines()
    split = next((i for i, l in enumerate(jd_lines) if PREFERRED_HEADING.match(l)), len(jd_lines))
    required_pos = _find_skills("\n".join(jd_lines[:split]))
    preferred_pos = {k: v for k, v in _find_skills("\n".join(jd_lines[split:])).items()
                     if k not in required_pos}
    required = sorted(required_pos, key=required_pos.get)    # JD order = rough importance
    preferred = sorted(preferred_pos, key=preferred_pos.get)
    have = set(_find_skills(SESSION["resume"]))

    # relevant experience: project / internship lines that mention a JD skill
    experience, section = [], ""
    for line in SESSION["resume"].splitlines():
        s = line.strip()
        if s and s.isupper():
            section = s
            continue
        if any(section.startswith(k) for k in EXPERIENCE_SECTIONS):
            hit = [k for k in _find_skills(s) if k in required_pos or k in preferred_pos]
            if hit:
                experience.append({"line": s.lstrip("-• ").strip(), "skills": hit})

    gap = {
        "required_skills": required, "matched_skills": [s for s in required if s in have],
        "missing_skills": [s for s in required if s not in have],
        "preferred_skills": preferred, "preferred_matched": [s for s in preferred if s in have],
        "preferred_missing": [s for s in preferred if s not in have],
        "relevant_experience": experience,
    }
    SESSION["gap"] = gap
    return _ok(f"{len(required)} required skills identified: {len(gap['matched_skills'])} matched, "
               f"{len(gap['missing_skills'])} missing ({', '.join(gap['missing_skills']) or 'none'}). "
               f"Preferred: {len(gap['preferred_matched'])}/{len(preferred)} matched.", **gap)


@tool
@safe_tool
def calculate_match_score() -> str:
    """Calculate the Resume-JD Skill Match Score = matched required skills / required skills x 100.
    It is a simple keyword-overlap indicator, NOT a hiring probability. Needs analyze_skill_gap first."""
    gap = SESSION.get("gap")
    if not gap:
        return _err("no skill-gap analysis available - call analyze_skill_gap first")
    total, matched = len(gap["required_skills"]), len(gap["matched_skills"])
    if total == 0:
        return _err("no recognisable required skills were found in the job description")
    score = round(100 * matched / total)
    band = "strong overlap" if score >= 75 else "moderate overlap" if score >= 50 else "limited overlap"
    SESSION["score"] = {"score": score, "matched": matched, "total": total, "band": band}
    return _ok(f"Resume-JD Skill Match Score: {score}% ({matched} of {total} required skills; {band}).",
               **SESSION["score"], formula=f"{matched} / {total} x 100")


@tool
@safe_tool
def generate_report(recommendations: list[str], learning_areas: list[str], summary: str) -> str:
    """Create the structured improvement report and save it to outputs/improvement_report.md.
    YOU write 'recommendations' (3-5 concrete, honest resume edits), 'learning_areas' (what to learn
    for the missing skills) and a 2-3 sentence 'summary'; the tool fills in score, skills and
    experience from working memory. Needs calculate_match_score first."""
    gap, sc = SESSION.get("gap"), SESSION.get("score")
    if not gap or not sc:
        return _err("analysis incomplete - run analyze_skill_gap and calculate_match_score first")
    bullets = lambda items: "\n".join(f"- {i}" for i in items) or "- None"
    exp = [f"- {e['line']}  _(shows: {', '.join(e['skills'])})_" for e in gap["relevant_experience"]]
    md = f"""# Resume-JD Analysis Report

## Resume-JD Skill Match Score
**{sc['score']}%** ({sc['matched']} of {sc['total']} required skills matched - {sc['band']}).
This is a keyword-overlap indicator for learning purposes, not a hiring prediction.

## Matched Skills
{bullets(gap['matched_skills'])}

## Missing Skills
{bullets(gap['missing_skills'])}

Preferred skills matched: {', '.join(gap['preferred_matched']) or 'none'}  |  missing: {', '.join(gap['preferred_missing']) or 'none'}

## Relevant Experience
{chr(10).join(exp) or '- No project/internship lines mention the required skills.'}

## Major Gaps
Top gaps, in the order the job description lists them:
{bullets(gap['missing_skills'][:3])}

## Recommended Resume Improvements
{bullets(recommendations)}

## Suggested Learning Areas
{bullets(learning_areas)}

## Overall Summary
{summary}
"""
    out = Path("outputs"); out.mkdir(exist_ok=True)
    (out / "improvement_report.md").write_text(md, encoding="utf-8")
    return _ok("Report generated and saved to outputs/improvement_report.md",
               path="outputs/improvement_report.md", report_markdown=md)


ALL_TOOLS = [load_documents, analyze_skill_gap, calculate_match_score, generate_report]
